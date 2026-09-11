import openpyxl
import os

FILE = "student_result.xlsx"


def menu():
    print("\n--- STUDENT RESULT MANAGEMENT SYSTEM ---")
    print("1. Get Result")
    print("2. View Previous Result")
    print("3. Exit")


def get_result():

    name = input("Enter student name: ")
    student_class = input("Enter class: ")

    math1 = int(input("Enter Math1 marks: "))
    chem = int(input("Enter Chemistry marks: "))
    cpro = int(input("Enter C Programming marks: "))
    mech = int(input("Enter Mechanics marks: "))
    cs = int(input("Enter CS marks: "))

    total = math1 + chem + cpro + mech + cs
    percentage = total / 5

    # New Excel file
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Result"

    # Headings
    ws.append([
        "Name", "Class", "Math1", "Chemistry",
        "C Programming", "Mechanics", "CS",
        "Total", "Percentage"
    ])

    # Result
    ws.append([
        name, student_class, math1, chem,
        cpro, mech, cs, total, percentage
    ])

    wb.save(FILE)

    print("\nResult saved successfully!")
    print("Total:", total)
    print("Percentage:", percentage)
    print("Excel file:", os.path.abspath(FILE))


def view_result():

    if not os.path.exists(FILE):
        print("\nNo previous result found.")
        return

    wb = openpyxl.load_workbook(FILE)
    ws = wb.active

    print("\n--- PREVIOUS RESULT ---")

    for row in ws.iter_rows(values_only=True):
        print(row)

    print("\nExcel file location:")
    print(os.path.abspath(FILE))


while True:

    menu()

    choice = int(input("Enter your choice: "))

    if choice == 1:
        get_result()

    elif choice == 2:
        view_result()

    elif choice == 3:
        print("Thank you!")
        break

    else:
        print("Invalid choice!")
        os.startfile(FILE_NAME)